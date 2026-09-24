import pandas as pd
from openpyxl import load_workbook
import os
from urllib.parse import urlparse, parse_qs
from datetime import date
import smtplib
from email.message import EmailMessage
from email.mime.text import MIMEText
import gspread
from google.auth import exceptions
from google.auth.transport.requests import Request
from google.oauth2.service_account import Credentials
import time

print('this will send emails of the video links')

def extract_real_url(google_url):
    if not isinstance(google_url, str):
        return None

    parsed = urlparse(google_url)
    params = parse_qs(parsed.query)

    return params.get("q", [None])[0]

os.chdir("/Users/chadgahafer/desktop/scripts")
print(os.getcwd())

file_path = "one church one vision reading plan.xlsx"

sheet_name = "Sheet1"

wb = load_workbook(file_path, data_only=True)
ws = wb[sheet_name]

# Find the column index of the hyperlink column
header = [cell.value for cell in ws[1]]

OT_link_col_idx = header.index("Old Testament Video") + 1  # 1-based index
NT_link_col_idx = header.index("New Testament Video") + 1  # 1-based index

OT_hyperlinks = []
NT_hyperlinks = []

df = pd.read_excel(file_path, sheet_name=sheet_name)

for excel_row in range(2, len(df) + 2):
    OT_cell = ws.cell(row=excel_row, column=OT_link_col_idx)
    if OT_cell.hyperlink and OT_cell.hyperlink.target:
        OT_link = extract_real_url(OT_cell.hyperlink.target)
        OT_hyperlinks.append(OT_link)
    else:
        OT_hyperlinks.append(None)
    
    NT_cell = ws.cell(row=excel_row, column=NT_link_col_idx)
    if NT_cell.hyperlink and NT_cell.hyperlink.target:
        NT_link = extract_real_url(NT_cell.hyperlink.target)
        NT_hyperlinks.append(NT_link)
    else:
        NT_hyperlinks.append(None)



df["OT Link"] = OT_hyperlinks
df["NT Link"] = NT_hyperlinks


df.to_excel('plan with links.xlsx', sheet_name='MyNewSheet', index=False)

TEST_DATE = None #using to make sure that the right week number gets pulled
# Example:
# TEST_DATE = date(2026, 1, 7)  # Wednesday


run_date = TEST_DATE if TEST_DATE else date.today()
iso_week_number = run_date.isocalendar()[1]

print(f"Run date: {run_date}")
print(f"ISO Week Number: {iso_week_number}")

current_week_df = df[df['Week'] == iso_week_number]

emailed_dataframe = current_week_df[['Dates',
                                    'Day Name',
                                    'Reading Day',
                                    'Old Testament Reading',
                                    'OT Link',
                                    'Psalm Reading', 
                                    'New Testament Reading',
                                    'NT Link']]

print(emailed_dataframe)

def send_email(subject, body, sender, recipients, password):
    msg = MIMEText(body, "html")
    msg['Subject'] = subject
    msg['From'] = sender
    msg['To'] = ', '.join(recipients)

    with smtplib.SMTP_SSL('smtp.gmail.com', 465) as smtp_server:
        smtp_server.login(sender, password)
        smtp_server.sendmail(sender, recipients, msg.as_string())

sender = 'chad.gahafer@gmail.com'
password = "nsnw ortu kcnn mjit"

subject = f"One Vision, One Story Reading Plan Week {iso_week_number}"
body = current_week_df.to_string(index=False)


scope = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]

# Provide the path to your service account JSON file
creds = Credentials.from_service_account_file(
    '/Users/chadgahafer/Desktop/Scripts/gspread-access-457201-9d4a7df64f34.json', scopes=scope)

# Authenticate and create the client
client = gspread.authorize(creds)

# Open the Google Sheet by title
sheet = client.open("2026 Bible Reading Plan Email List").worksheet("Email List") # Or use sheet index/URL

email_sheet_values = sheet.get_all_values()

recipients = [
    row[0].strip()
    for row in email_sheet_values
    if row and row[0].strip() and "@" in row[0]
]

html_table = emailed_dataframe.to_html(
    index=False,
    escape=False,
    border=0,
    justify="center"
)

html_table = (
    html_table
    .replace("<th>", "<th style='border:1px solid #ccc; padding:6px; background:#f2f2f2; text-align:center;'>")
    .replace("<td>", "<td style='border:1px solid #ccc; padding:6px; text-align:center;'>")
)


html_body = f"""
<html>
  <body style="font-family: Arial, sans-serif;">
    <p style="font-size:14px;">Here is this week's reading plan:</p>

    <table style="
        border-collapse: collapse;
        width: 100%;
        max-width: 900px;
        margin: auto;
        font-size: 13px;
        ">
        {html_table[html_table.find('<thead>'):]}
    </table>

    <p style="font-size:14px;">You received this email because you elected to get weekly emails with the reading plan. If you'd like to be removed from the list please email me at chad.gahafer@gmail.com and I will take care of that for you.</p>
  </body>
</html>
"""

# Logic automatically delay email sending rate based on amount of emails
recipient_count = len(recipients)

if recipient_count <= 30:
    delay = 0
elif recipient_count <= 75:
    delay = 0.5
elif recipient_count <= 150:
    delay = 1
else:
    delay = 2

print(f"Sending to {recipient_count} recipients")


for email in recipients:
    try:
        send_email(subject, html_body, sender, [email], password)
        print(f'Email successfully sent to {email}')
        if delay:
            time.sleep(delay)
    except Exception as e:
        print(f'Email FAILURE for {email}: {e}')